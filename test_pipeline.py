"""
test_pipeline.py
================
Pytest suite for pipeline.py integration glue.

Run:
    pytest test_pipeline.py -v

Tests cover:
  · Full pipeline on .docx (no client — citation step skipped)
  · Full pipeline on .txt
  · Pipeline fails gracefully on bad input file
  · Citation step skipped when skip_citations=True
  · Citation step skipped when client is None
  · Citation reformatted when citations.py present (mocked)
  · Output file exists and is valid .docx
  · PipelineResult.report() renders correctly
  · Batch pipeline: multiple files
  · Batch pipeline: recovers from one bad file
  · Change log populated
  · Warnings propagated from layout step
"""

import os
import sys
import tempfile
import pytest
from unittest.mock import MagicMock, patch
from pathlib import Path

from docx import Document

from pipeline import (
    run_pipeline, run_batch_pipeline,
    PipelineResult, CitationResult,
    _citations_available, _try_reformat_citations,
)
from layout_formatter import FormattedResult
from style_profiles import get_profile


# ─────────────────────────────────────────────────────────────────────
# FIXTURES
# ─────────────────────────────────────────────────────────────────────

@pytest.fixture
def tmp_dir():
    with tempfile.TemporaryDirectory() as d:
        yield d


def _make_docx(tmp_dir: str, name: str = "manuscript.docx") -> str:
    doc  = Document()
    doc.add_paragraph("A Study on Network Intrusion Detection")
    doc.add_paragraph(
        "Abstract: This paper presents a comprehensive analysis of deep learning "
        "methods for intrusion detection on the CICIDS-2017 dataset."
    )
    doc.add_paragraph("Introduction")
    doc.add_paragraph(
        "Network security is critical in modern infrastructure. "
        "Recent advances in machine learning have enabled highly accurate "
        "detection of malicious traffic patterns [1]."
    )
    doc.add_paragraph("Methods")
    doc.add_paragraph(
        "We trained an LSTM on 80% of the dataset and evaluated on 20%. "
        "Hyperparameters were tuned via grid search."
    )
    doc.add_paragraph("Results")
    doc.add_paragraph("The model achieved 98.7% accuracy and F1=0.97.")
    doc.add_paragraph("References")
    doc.add_paragraph("[1] Smith, J. et al. Deep learning for IDS. IEEE, 2022.")
    path = os.path.join(tmp_dir, name)
    doc.save(path)
    return path


def _make_txt(tmp_dir: str) -> str:
    content = (
        "A STUDY ON NETWORK INTRUSION DETECTION\n\n"
        "Abstract: This paper presents a deep learning approach.\n\n"
        "INTRODUCTION\n\n"
        "Background text here.\n\n"
        "REFERENCES\n\n"
        "[1] Smith, J. et al. 2022."
    )
    path = os.path.join(tmp_dir, "manuscript.txt")
    with open(path, "w") as f:
        f.write(content)
    return path


# ─────────────────────────────────────────────────────────────────────
# 1. BASIC PIPELINE RUNS
# ─────────────────────────────────────────────────────────────────────

class TestBasicPipeline:

    def test_pipeline_docx_ieee_no_client(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        out    = os.path.join(tmp_dir, "out.docx")
        result = run_pipeline(path, "ieee", out, anthropic_client=None)

        assert result.success
        assert os.path.exists(out)
        assert result.profile.name == "ieee"

    def test_pipeline_docx_apa7(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        out    = os.path.join(tmp_dir, "out.docx")
        result = run_pipeline(path, "apa7", out)

        assert result.success
        assert os.path.exists(out)

    def test_pipeline_docx_harvard(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        out    = os.path.join(tmp_dir, "out.docx")
        result = run_pipeline(path, "harvard", out)

        assert result.success
        assert os.path.exists(out)

    def test_pipeline_txt_input(self, tmp_dir):
        path   = _make_txt(tmp_dir)
        out    = os.path.join(tmp_dir, "out.docx")
        result = run_pipeline(path, "ieee", out)

        assert result.success
        assert os.path.exists(out)
        # Output must be a valid docx
        doc = Document(out)
        assert len(doc.paragraphs) > 0

    def test_pipeline_auto_names_output(self, tmp_dir):
        path   = _make_docx(tmp_dir, "mypaper.docx")
        result = run_pipeline(path, "ieee")

        expected = os.path.join(tmp_dir, "mypaper_ieee.docx")
        assert result.output_path == expected
        assert os.path.exists(expected)

    def test_output_is_valid_docx(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        out    = os.path.join(tmp_dir, "out.docx")
        run_pipeline(path, "ieee", out)
        doc = Document(out)
        assert len(doc.paragraphs) > 0


# ─────────────────────────────────────────────────────────────────────
# 2. CITATION STEP BEHAVIOUR
# ─────────────────────────────────────────────────────────────────────

class TestCitationStep:

    def test_citations_skipped_when_no_client(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        out    = os.path.join(tmp_dir, "out.docx")
        result = run_pipeline(path, "ieee", out, anthropic_client=None)

        assert result.citation_result.skipped is True

    def test_citations_skipped_when_flag_set(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        out    = os.path.join(tmp_dir, "out.docx")
        mock   = MagicMock()   # client provided but skip=True
        result = run_pipeline(path, "ieee", out,
                              anthropic_client=mock,
                              skip_citations=True)

        assert result.citation_result.skipped is True

    def test_citations_fallback_when_module_missing(self, tmp_dir):
        """citations.py not importable → graceful skip with warning."""
        path   = _make_docx(tmp_dir)
        out    = os.path.join(tmp_dir, "out.docx")
        mock   = MagicMock()

        # Ensure citations is not importable
        with patch.dict(sys.modules, {"citations": None}):
            result = _try_reformat_citations(
                Document(), "ieee", mock, log_fn=lambda m: None
            )

        assert result.skipped is True
        assert any("citations.py" in w for w in result.warnings)

    def test_citations_called_when_module_present(self, tmp_dir):
        """Mock citations.py present → reformat_citations is called."""
        path   = _make_docx(tmp_dir)
        out    = os.path.join(tmp_dir, "out.docx")
        mock_client = MagicMock()

        # Build a fake citations module
        fake_citations = MagicMock()
        fake_style = MagicMock()
        fake_style.name = "APA"
        fake_citations.detect_style.return_value = fake_style

        fake_result = MagicMock()
        fake_result.warnings      = []
        fake_result.changed_count = 3
        fake_citations.reformat_citations.return_value = fake_result

        with patch.dict(sys.modules, {"citations": fake_citations}):
            cit_result = _try_reformat_citations(
                Document(), "ieee", mock_client, log_fn=lambda m: None
            )

        fake_citations.reformat_citations.assert_called_once()
        assert cit_result.changed_count == 3
        assert not cit_result.skipped

    def test_citations_already_correct_style_skips(self, tmp_dir):
        """If source style == target style, reformat is skipped."""
        fake_citations = MagicMock()
        fake_style     = MagicMock()
        fake_style.name = "ieee"
        fake_citations.detect_style.return_value = fake_style

        with patch.dict(sys.modules, {"citations": fake_citations}):
            cit_result = _try_reformat_citations(
                Document(), "ieee", MagicMock(), log_fn=lambda m: None
            )

        assert cit_result.skipped is True
        fake_citations.reformat_citations.assert_not_called()

    def test_citation_error_produces_warning_not_crash(self, tmp_dir):
        """If reformat_citations raises, pipeline continues with warning."""
        fake_citations = MagicMock()
        fake_style     = MagicMock()
        fake_style.name = "apa"
        fake_citations.detect_style.return_value = fake_style
        fake_citations.reformat_citations.side_effect = RuntimeError("API error")

        with patch.dict(sys.modules, {"citations": fake_citations}):
            cit_result = _try_reformat_citations(
                Document(), "ieee", MagicMock(), log_fn=lambda m: None
            )

        assert cit_result.skipped is True
        assert any("error" in w.lower() for w in cit_result.warnings)


# ─────────────────────────────────────────────────────────────────────
# 3. FAILURE MODES
# ─────────────────────────────────────────────────────────────────────

class TestFailureModes:

    def test_bad_input_path_fails_gracefully(self, tmp_dir):
        out    = os.path.join(tmp_dir, "out.docx")
        result = run_pipeline("/nonexistent/file.docx", "ieee", out)

        assert result.success is False
        assert len(result.errors) > 0

    def test_unsupported_format_fails_gracefully(self, tmp_dir):
        bad = os.path.join(tmp_dir, "file.pdf")
        with open(bad, "w") as f:
            f.write("dummy")
        out    = os.path.join(tmp_dir, "out.docx")
        result = run_pipeline(bad, "ieee", out)

        assert result.success is False
        assert len(result.errors) > 0

    def test_unknown_style_raises(self, tmp_dir):
        path = _make_docx(tmp_dir)
        with pytest.raises(KeyError, match="not found"):
            run_pipeline(path, "nonexistent_style")

    def test_result_has_errors_list_on_failure(self, tmp_dir):
        result = run_pipeline("/bad/path.docx", "ieee")
        assert isinstance(result.errors, list)
        assert len(result.errors) > 0


# ─────────────────────────────────────────────────────────────────────
# 4. PIPELINE RESULT
# ─────────────────────────────────────────────────────────────────────

class TestPipelineResult:

    def test_changes_populated(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        result = run_pipeline(path, "ieee")

        assert len(result.all_changes) > 0

    def test_warnings_propagated_from_layout(self, tmp_dir):
        """Long abstract should trigger a word-limit warning."""
        doc = Document()
        doc.add_paragraph("Short title")
        doc.add_paragraph("Abstract: " + "word " * 250)  # over IEEE 200-word limit
        doc.add_paragraph("Body text here.")
        path = os.path.join(tmp_dir, "long_abstract.docx")
        doc.save(path)

        result = run_pipeline(path, "ieee")
        assert any("Abstract" in w for w in result.warnings), \
            f"Expected abstract warning, got: {result.warnings}"

    def test_report_renders(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        result = run_pipeline(path, "ieee")
        report = result.report()

        assert "Pipeline Result" in report
        assert "ieee" in report.lower() or "IEEE" in report

    def test_change_summary_string(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        result = run_pipeline(path, "ieee")

        summary = result.change_summary
        assert "layout changes" in summary

    def test_elapsed_time_positive(self, tmp_dir):
        path   = _make_docx(tmp_dir)
        result = run_pipeline(path, "ieee")
        assert result.elapsed_sec > 0


# ─────────────────────────────────────────────────────────────────────
# 5. BATCH PIPELINE
# ─────────────────────────────────────────────────────────────────────

class TestBatchPipeline:

    def test_batch_all_succeed(self, tmp_dir):
        paths = [
            _make_docx(tmp_dir, "doc1.docx"),
            _make_docx(tmp_dir, "doc2.docx"),
        ]
        out_dir = os.path.join(tmp_dir, "output")
        results = run_batch_pipeline(paths, "ieee", output_dir=out_dir)

        assert len(results) == 2
        assert all(r.success for r in results)
        assert os.path.exists(os.path.join(out_dir, "doc1_ieee.docx"))
        assert os.path.exists(os.path.join(out_dir, "doc2_ieee.docx"))

    def test_batch_recovers_from_bad_file(self, tmp_dir):
        good = _make_docx(tmp_dir, "good.docx")
        bad  = os.path.join(tmp_dir, "bad.xyz")
        with open(bad, "w") as f:
            f.write("not a document")

        results = run_batch_pipeline([good, bad], "ieee",
                                     output_dir=tmp_dir,
                                     log_fn=lambda m: None)

        assert len(results) == 2
        good_r = next(r for r in results if "good" in r.output_path)
        bad_r  = next(r for r in results if "bad"  in r.output_path)
        assert good_r.success is True
        assert bad_r.success  is False

    def test_batch_mixed_formats(self, tmp_dir):
        docx_path = _make_docx(tmp_dir)
        txt_path  = _make_txt(tmp_dir)
        out_dir   = os.path.join(tmp_dir, "out")

        results = run_batch_pipeline([docx_path, txt_path], "harvard",
                                     output_dir=out_dir,
                                     log_fn=lambda m: None)
        assert len(results) == 2
        assert all(r.success for r in results)

    def test_batch_returns_one_result_per_input(self, tmp_dir):
        paths   = [_make_docx(tmp_dir, f"doc{i}.docx") for i in range(3)]
        results = run_batch_pipeline(paths, "apa7",
                                     log_fn=lambda m: None)
        assert len(results) == 3


# ─────────────────────────────────────────────────────────────────────
# 6. CITATIONS AVAILABLE HELPER
# ─────────────────────────────────────────────────────────────────────

class TestCitationsAvailable:

    def test_returns_false_when_not_installed(self):
        with patch.dict(sys.modules, {"citations": None}):
            assert _citations_available() is False

    def test_returns_true_when_module_present(self):
        fake = MagicMock()
        with patch.dict(sys.modules, {"citations": fake}):
            assert _citations_available() is True
